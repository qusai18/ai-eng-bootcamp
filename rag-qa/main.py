import io
import json
import math
import os
import re
import zipfile
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Literal
from xml.etree import ElementTree

import numpy as np
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from openai import OpenAI, OpenAIError
from pydantic import BaseModel, Field

from chunking import GUIDE, STRATEGIES, chunk_documents

ROOT = Path(__file__).resolve().parent
_default_ratings = "/tmp/rag-qa-ratings.json" if os.name != "nt" else str(ROOT / "ratings.json")
RATINGS_PATH = Path(os.getenv("RATINGS_PATH", _default_ratings))
load_dotenv(ROOT.parent / ".env")
load_dotenv(ROOT / ".env")

EMBED_MODEL = "text-embedding-3-small"
EMBED_USD_PER_MILLION = 0.02
CHAT_MODEL = "gpt-4o-mini"
TOP_K = 3
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARS = 40000
ALLOWED_SUFFIXES = {".txt", ".pdf", ".docx", ".ppt", ".pptx"}
ANSWER_PROMPT = (
    "Answer the question using ONLY the context below. "
    "If the context doesn't contain the answer, say so.\n\n"
    "Context:\n{context}\n\n"
    "Question: {question}"
)


class Index:
    def __init__(self, chunks: list[str], vectors: np.ndarray, tokens: int) -> None:
        self.chunks = chunks
        self.vectors = vectors
        self.tokens = tokens


class Store:
    def __init__(self) -> None:
        self.mode = "single"
        self.documents: list[tuple[str, str]] = []
        self.indexes: dict[str, Index] = {}
        self.ratings: list[dict[str, str | int]] = _load_ratings()
        self.report: dict | None = None


def _load_ratings() -> list[dict[str, str | int]]:
    if not RATINGS_PATH.exists():
        return []
    try:
        data = json.loads(RATINGS_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    if not isinstance(data, list):
        return []
    return data


store = Store()


def _save_ratings() -> None:
    RATINGS_PATH.parent.mkdir(parents=True, exist_ok=True)
    RATINGS_PATH.write_text(json.dumps(store.ratings), encoding="utf-8")


def _client() -> OpenAI:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise HTTPException(status_code=500, detail="OPENAI_API_KEY is not set")
    return OpenAI(api_key=api_key)


def _embed(texts: list[str]) -> tuple[np.ndarray, int]:
    response = _client().embeddings.create(model=EMBED_MODEL, input=texts)
    ordered = sorted(response.data, key=lambda item: item.index)
    vectors = np.array([item.embedding for item in ordered], dtype=np.float32)
    norms = np.linalg.norm(vectors, axis=1, keepdims=True)
    norms[norms == 0] = 1
    tokens = int(response.usage.total_tokens or 0) if response.usage else 0
    return vectors / norms, tokens


def _chunk_performance(chunks: list[str]) -> float:
    """Score how usable the chunks are, from 0 to 100.

    A chunk scores well when its text is in the 200-1200 character range
    and the average length stays near 600 characters.
    """
    if not chunks:
        return 0.0
    lengths: list[int] = []
    for chunk in chunks:
        body = chunk.split("\n", 1)[-1] if chunk.startswith("Source:") else chunk
        lengths.append(len(body.strip()))
    in_range = sum(1 for length in lengths if 200 <= length <= 1200) / len(lengths)
    average = sum(lengths) / len(lengths)
    centered = max(0.0, 1 - abs(average - 600) / 800)
    return round(100 * (0.65 * in_range + 0.35 * centered), 1)


def _cost(tokens: int) -> float:
    return round(tokens * EMBED_USD_PER_MILLION / 1_000_000, 8)


def _require_strategy(strategy: str) -> str:
    if strategy not in STRATEGIES:
        raise HTTPException(status_code=400, detail=f"Unknown strategy: {strategy}")
    return strategy


def ensure_index(strategy: str) -> Index:
    _require_strategy(strategy)
    if not store.documents:
        raise HTTPException(status_code=400, detail="Add a document before asking")
    cached = store.indexes.get(strategy)
    if cached is not None:
        return cached

    spent = 0

    def embed_sentences(sentences: list[str]) -> np.ndarray:
        nonlocal spent
        vectors, tokens = _embed(sentences)
        spent += tokens
        return vectors

    try:
        chunks = chunk_documents(
            store.documents,
            strategy,
            store.mode,
            embed_sentences if strategy == "semantic" else None,
        )
    except OpenAIError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    if not chunks:
        raise HTTPException(status_code=400, detail="That document produced no chunks")
    try:
        vectors, tokens = _embed(chunks)
        spent += tokens
    except OpenAIError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    index = Index(chunks, vectors, spent)
    store.indexes[strategy] = index
    return index


def retrieve(question: str, strategy: str) -> list[str]:
    index = ensure_index(strategy)
    try:
        query_vectors, _tokens = _embed([question])
        query = query_vectors[0]
    except OpenAIError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    scores = index.vectors @ query
    order = np.argsort(scores)[::-1][:TOP_K]
    return [index.chunks[int(i)] for i in order]


def answer_question(question: str, strategy: str) -> tuple[str, list[str]]:
    sources = retrieve(question, strategy)
    prompt = ANSWER_PROMPT.format(context="\n\n".join(sources), question=question)
    try:
        completion = _client().chat.completions.create(
            model=CHAT_MODEL,
            temperature=0,
            messages=[{"role": "user", "content": prompt}],
        )
    except OpenAIError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    return completion.choices[0].message.content or "", sources


def scoreboard() -> list[dict[str, str | float | int]]:
    totals: dict[str, list[int]] = {key: [] for key in STRATEGIES}
    for rating in store.ratings:
        strategy = str(rating["strategy"])
        if strategy in totals:
            totals[strategy].append(int(rating["score"]))
    rows = []
    for key, scores in totals.items():
        rows.append(
            {
                "strategy": key,
                "name": STRATEGIES[key]["name"],
                "ratings": len(scores),
                "average": round(sum(scores) / len(scores), 2) if scores else 0,
            }
        )
    return rows


def recommendation() -> dict[str, str | float | int | None]:
    rows = [row for row in scoreboard() if int(row["ratings"]) > 0]
    guide = GUIDE[store.mode]
    suggested = "document" if store.mode == "multiple" else "recursive"
    if not rows:
        return {
            "strategy": None,
            "name": None,
            "average": None,
            "ratings": 0,
            "guide_strategy": suggested,
            "guide_name": STRATEGIES[suggested]["name"],
            "reason": guide,
        }
    best = max(rows, key=lambda row: (float(row["average"]), int(row["ratings"])))
    return {
        "strategy": best["strategy"],
        "name": best["name"],
        "average": best["average"],
        "ratings": best["ratings"],
        "guide_strategy": suggested,
        "guide_name": STRATEGIES[suggested]["name"],
        "reason": (
            f"{best['name']} has the highest score from your ratings "
            f"({best['average']} / 5 across {best['ratings']} rating"
            f"{'' if best['ratings'] == 1 else 's'}). {guide}"
        ),
    }


@asynccontextmanager
async def lifespan(_app: FastAPI):
    yield


app = FastAPI(lifespan=lifespan)


class AskRequest(BaseModel):
    question: str = Field(min_length=1, max_length=2000)
    strategy: str


class CompareRequest(BaseModel):
    question: str = Field(min_length=1, max_length=2000)


class RateRequest(BaseModel):
    strategy: str
    question: str = Field(min_length=1, max_length=2000)
    score: int = Field(ge=0, le=5)


class AskResponse(BaseModel):
    answer: str
    sources: list[str]
    model: str
    strategy: str
    chunk_count: int


@app.get("/health")
def health() -> dict[str, str | int]:
    ready = store.indexes.get("recursive")
    return {"status": "ok", "chunks": 0 if ready is None else len(ready.chunks)}


@app.get("/")
def home() -> FileResponse:
    return FileResponse(ROOT / "index.html")


@app.get("/sample")
def sample() -> dict[str, str]:
    text = (ROOT / "sample_rag_document.txt").read_text(encoding="utf-8")
    return {"name": "sample_rag_document.txt", "text": text}


@app.get("/strategies")
def strategies() -> dict:
    return {
        "strategies": [
            {"id": key, **value, "suggested": key == ("document" if store.mode == "multiple" else "recursive")}
            for key, value in STRATEGIES.items()
        ],
        "mode": store.mode,
        "documents": [name for name, _text in store.documents],
        "scores": scoreboard(),
        "recommendation": recommendation(),
    }


def _docx_page_count(data: bytes) -> int | None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "docProps/app.xml" not in archive.namelist():
                return None
            root = ElementTree.fromstring(archive.read("docProps/app.xml"))
    except (zipfile.BadZipFile, ElementTree.ParseError, KeyError):
        return None
    for node in root.iter():
        if node.tag.endswith("Pages") and node.text and node.text.strip().isdigit():
            pages = int(node.text.strip())
            return pages or None
    return None


def _display_title(raw: str | None, filename: str, text: str) -> str:
    title = re.sub(r"\s+", " ", (raw or "")).strip(" \t\r\n")
    if title and title.lower() not in {"untitled", "unknown"}:
        return title[:120]
    for line in text.splitlines():
        cleaned = line.strip().lstrip("#").strip()
        if 3 <= len(cleaned) <= 90:
            return cleaned
    stem = Path(filename).stem.replace("_", " ").replace("-", " ").strip()
    return stem or filename


def _structure(text: str) -> tuple[int, int, float]:
    headings = 0
    lists = 0
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for line in lines:
        if line.startswith("#") or (len(line) <= 60 and line.isupper() and any(char.isalpha() for char in line)):
            headings += 1
        elif re.match(r"^(?:[-*•]|\d+[.)])\s+\S", line):
            lists += 1
    ratio = lists / len(lines) if lines else 0.0
    return headings, lists, ratio


def _read_document(name: str, data: bytes) -> dict:
    suffix = Path(name).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise HTTPException(
            status_code=400,
            detail=f"{name} must be a TXT, PDF, DOCX, PPT, or PPTX file",
        )
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=400, detail=f"{name} is larger than 10 MB")
    title = ""
    pages = 1
    pages_estimated = False
    try:
        if suffix == ".txt":
            text = data.decode("utf-8", errors="replace")
            pages_estimated = True
        elif suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            pages = max(len(reader.pages), 1)
            meta = reader.metadata
            title = str(getattr(meta, "title", "") or "")
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            from docx import Document

            document = Document(io.BytesIO(data))
            title = str(document.core_properties.title or "")
            parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n".join(parts)
            saved_pages = _docx_page_count(data)
            if saved_pages:
                pages = saved_pages
            else:
                pages_estimated = True
        else:
            if suffix == ".ppt" and not data.startswith(b"PK"):
                raise HTTPException(
                    status_code=400,
                    detail=f"{name} is an older PowerPoint file. Save it as .pptx and upload that.",
                )
            from pptx import Presentation

            presentation = Presentation(io.BytesIO(data))
            title = str(presentation.core_properties.title or "")
            pages = max(len(presentation.slides), 1)
            parts: list[str] = []
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                        parts.append(shape.text_frame.text.strip())
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        parts.append(notes)
            text = "\n\n".join(parts)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read {name}: {exc}") from exc

    text = text.strip()
    if not text:
        raise HTTPException(status_code=400, detail=f"No text found in {name}")
    text = text[:MAX_EXTRACTED_CHARS]
    words = len(re.findall(r"\b[\w']+\b", text))
    if pages_estimated:
        pages = max(1, math.ceil(words / 400)) if words else 1
    headings, _lists, list_ratio = _structure(text)
    return {
        "filename": name,
        "title": _display_title(title, name, text),
        "pages": pages,
        "pages_estimated": pages_estimated,
        "words": words,
        "characters": len(text),
        "headings": headings,
        "list_ratio": list_ratio,
        "text": text,
    }


def _extract_text(name: str, data: bytes) -> str:
    return _read_document(name, data)["text"]


def recommend_strategy(profiles: list[dict]) -> dict[str, str]:
    if len(profiles) > 1:
        key = "document"
        reason = "Several files are uploaded, so document-based chunking keeps each file as its own source."
    else:
        doc = profiles[0]
        words = int(doc["words"])
        headings = int(doc["headings"])
        if words < 180 and headings < 2:
            key = "fixed"
            reason = "This is a short note, so fixed-size chunks are enough and cheaper to embed."
        elif float(doc["list_ratio"]) >= 0.3 and words >= 120:
            key = "adaptive"
            reason = "The text mixes lists and prose, so adaptive chunking keeps dense lines smaller than the surrounding paragraphs."
        elif headings >= 3:
            key = "hierarchical"
            reason = "The document is organized into sections, so hierarchical chunking can store each section overview with its details."
        elif headings < 2 and words >= 700:
            key = "semantic"
            reason = "The document is long and has few headings, so semantic chunking can split where the topic changes."
        else:
            key = "recursive"
            reason = "The document reads like continuous prose, so recursive chunking can keep paragraphs together."
    return {"strategy": key, "name": STRATEGIES[key]["name"], "reason": reason}


def _fallback_prompts(documents: list[tuple[str, str]]) -> list[dict[str, str]]:
    prompts: list[dict[str, str]] = []
    for _name, text in documents:
        for raw in text.splitlines():
            line = raw.strip().lstrip("#").strip(" -")
            if len(line) < 12 or len(line) > 90:
                continue
            label = line if len(line) <= 32 else f"{line[:29].rstrip()}…"
            prompts.append({"label": label, "question": f"What does the document say about {line}?"})
            if len(prompts) == 3:
                return prompts
    if not documents:
        return []
    title = Path(documents[0][0]).stem.replace("_", " ")
    return [{"label": title[:32] or "Document", "question": f"What is {title or 'this document'} about?"}]


def _fallback_summary(text: str) -> str:
    sentences = [part.strip() for part in re.split(r"(?<=[.!?])\s+", text) if part.strip()]
    summary = " ".join(sentences[:2]).strip()
    return summary[:500] or "The file contains text, but it has no complete sentences to summarize."


def _model_notes(documents: list[tuple[str, str]]) -> tuple[list[dict[str, str]], dict[str, str]]:
    excerpts = "\n\n".join(f"# {name}\n{text[:3500]}" for name, text in documents)[:9000]
    instruction = (
        "Read the documents. Return JSON with two keys. "
        '"prompts": exactly 3 objects {"label":"2 to 4 words","question":"one sentence"} '
        "that a reader could answer from the text. "
        '"synopses": one object per document {"name":"the filename","summary":"two sentences about that document"}. '
        "Use only facts from the text.\n\n"
        f"{excerpts}"
    )
    try:
        completion = _client().chat.completions.create(
            model=CHAT_MODEL,
            temperature=0,
            response_format={"type": "json_object"},
            messages=[{"role": "user", "content": instruction}],
        )
        payload = json.loads(completion.choices[0].message.content or "{}")
    except (OpenAIError, json.JSONDecodeError, AttributeError, TypeError):
        return _fallback_prompts(documents), {}
    if not isinstance(payload, dict):
        return _fallback_prompts(documents), {}
    prompts: list[dict[str, str]] = []
    raw = payload.get("prompts")
    if isinstance(raw, list):
        for item in raw:
            if not isinstance(item, dict):
                continue
            question = str(item.get("question") or "").strip()
            label = str(item.get("label") or "").strip()
            if not question:
                continue
            prompts.append({"label": (label or question)[:40], "question": question[:300]})
            if len(prompts) == 3:
                break
    summaries: dict[str, str] = {}
    raw_synopses = payload.get("synopses")
    if isinstance(raw_synopses, list):
        for item in raw_synopses:
            if not isinstance(item, dict):
                continue
            filename = str(item.get("name") or "").strip()
            summary = str(item.get("summary") or "").strip()
            if filename and summary:
                summaries[filename] = summary[:500]
    return prompts or _fallback_prompts(documents), summaries


def describe_documents(profiles: list[dict]) -> dict:
    documents = [(str(profile["filename"]), str(profile["text"])) for profile in profiles]
    prompts, summaries = _model_notes(documents)
    public = []
    for profile in profiles:
        filename = str(profile["filename"])
        public.append(
            {
                "filename": filename,
                "title": profile["title"],
                "pages": profile["pages"],
                "pages_estimated": profile["pages_estimated"],
                "words": profile["words"],
                "characters": profile["characters"],
                "summary": summaries.get(filename) or _fallback_summary(str(profile["text"])),
            }
        )
    return {
        "prompts": prompts,
        "documents": public,
        "recommendation": recommend_strategy(profiles),
    }


def _index_documents(mode: str, strategy: str, documents: list[tuple[str, str]]) -> dict:
    if mode == "single" and len(documents) != 1:
        raise HTTPException(status_code=400, detail="Single document accepts one file")
    if not documents or len(documents) > 8:
        raise HTTPException(status_code=400, detail="Upload between 1 and 8 files")
    store.mode = mode
    store.documents = documents
    store.indexes.clear()
    rows = []
    for key in STRATEGIES:
        index = ensure_index(key)
        rows.append(
            {
                "strategy": key,
                "name": STRATEGIES[key]["name"],
                "chunks": len(index.chunks),
                "tokens": index.tokens,
                "cost_usd": _cost(index.tokens),
                "performance": _chunk_performance(index.chunks),
            }
        )
    best = max(rows, key=lambda row: (float(row["performance"]), -int(row["tokens"])))
    total_tokens = sum(int(row["tokens"]) for row in rows)
    report = {
        "model": EMBED_MODEL,
        "price_per_million": EMBED_USD_PER_MILLION,
        "rows": rows,
        "total_tokens": total_tokens,
        "total_cost_usd": _cost(total_tokens),
        "best_strategy": best["strategy"],
        "best_name": best["name"],
        "best_performance": best["performance"],
        "best_tokens": best["tokens"],
        "best_cost_usd": best["cost_usd"],
    }
    store.report = report
    chosen = store.indexes[strategy]
    return {
        "mode": store.mode,
        "strategy": strategy,
        "documents": [name for name, _text in store.documents],
        "chunk_count": len(chosen.chunks),
        "chunks": [chunk[:280] for chunk in chosen.chunks[:8]],
        "recommendation": recommendation(),
        "report": report,
    }


@app.post("/reset")
def reset() -> dict[str, str]:
    store.mode = "single"
    store.documents = []
    store.indexes.clear()
    store.ratings = []
    store.report = None
    _save_ratings()
    return {"status": "cleared"}


@app.post("/prompts")
async def prompts(files: list[UploadFile] = File(...)) -> dict:
    if not files or len(files) > 8:
        raise HTTPException(status_code=400, detail="Upload between 1 and 8 files")
    profiles: list[dict] = []
    for upload in files:
        name = Path(upload.filename or "document").name
        data = await upload.read()
        profiles.append(_read_document(name, data))
    return describe_documents(profiles)


@app.post("/documents")
async def ingest(
    mode: Literal["single", "multiple"] = Form(...),
    strategy: str = Form(...),
    files: list[UploadFile] = File(...),
) -> dict:
    chosen = _require_strategy(strategy)
    if mode == "single" and len(files) != 1:
        raise HTTPException(status_code=400, detail="Single document accepts one file")
    documents: list[tuple[str, str]] = []
    for upload in files:
        name = Path(upload.filename or "document").name
        data = await upload.read()
        documents.append((name, _extract_text(name, data)))
    return _index_documents(mode, chosen, documents)


@app.post("/ask", response_model=AskResponse)
def ask(body: AskRequest) -> AskResponse:
    strategy = _require_strategy(body.strategy)
    question = body.question.strip()
    answer, sources = answer_question(question, strategy)
    index = store.indexes[strategy]
    return AskResponse(
        answer=answer,
        sources=sources,
        model=CHAT_MODEL,
        strategy=strategy,
        chunk_count=len(index.chunks),
    )


@app.post("/compare")
def compare(body: CompareRequest) -> dict:
    question = body.question.strip()
    results = []
    for strategy in STRATEGIES:
        answer, sources = answer_question(question, strategy)
        results.append(
            {
                "strategy": strategy,
                "name": STRATEGIES[strategy]["name"],
                "answer": answer,
                "sources": sources,
                "chunk_count": len(store.indexes[strategy].chunks),
            }
        )
    return {"question": question, "results": results, "recommendation": recommendation()}


@app.post("/rate")
def rate(body: RateRequest) -> dict:
    strategy = _require_strategy(body.strategy)
    store.ratings.append(
        {"strategy": strategy, "question": body.question.strip(), "score": body.score}
    )
    _save_ratings()
    return {"scores": scoreboard(), "recommendation": recommendation()}
